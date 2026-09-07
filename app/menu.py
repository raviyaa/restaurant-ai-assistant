from pathlib import Path

from docx import Document as WordDocument
from langchain_core.documents import Document

CATEGORIES = [
    "Simply Pizza Range",
    "Classic Pizza Range",
    "Supreme Pizza Range",
    "Premium Pizza Range",
    "Chef's Special Pizza Range",
    "Vito's Special - XL Size Pizza",
    "Extra Toppings for Pizza",
    "Calzone",
    "Salad",
    "Pasta",
    "Lasanga",
    "Dessert",
    "Drinks",
]
SKIP_PREFIXES = ("Size:", "Choice of", "Folded", "Vito's Restaurant Menu")


def load_menu_docs(menu_path: Path) -> list[Document]:
    word_doc = WordDocument(menu_path)
    lines = [p.text.strip() for p in word_doc.paragraphs if p.text.strip()]

    docs: list[Document] = []
    category = None
    range_price = None
    notes: list[str] = []
    i = 0

    while i < len(lines):
        line = lines[i]
        if line in CATEGORIES:
            category, range_price, notes = line, None, []
            i += 1
            continue
        if line.startswith("Range Price:"):
            range_price = line.replace("Range Price:", "").strip()
            i += 1
            continue
        if line.startswith(SKIP_PREFIXES):
            notes.append(line)
            i += 1
            continue

        name = line
        price, ingredients = range_price, None
        i += 1
        while i < len(lines) and lines[i] not in CATEGORIES:
            if lines[i].startswith("Price:"):
                price = lines[i].replace("Price:", "").strip()
                i += 1
            elif lines[i].startswith("Ingredients/Description:"):
                ingredients = lines[i].replace("Ingredients/Description:", "").strip()
                i += 1
            else:
                break

        parts = [f"Category: {category}", f"Item: {name}"]
        if price:
            parts.append(f"Price: {price}")
        if ingredients:
            parts.append(f"Ingredients: {ingredients}")
        if notes:
            parts.append("Notes: " + " ".join(notes))

        docs.append(
            Document(
                page_content="\n".join(parts),
                metadata={"category": category or "", "item": name},
            )
        )

    return docs
